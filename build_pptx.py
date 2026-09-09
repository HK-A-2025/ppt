#!/usr/bin/env python3
"""
Generator Berkas Microsoft PowerPoint (.PPTX)
Makalah: Sejarah dan Perkembangan Hukum Perdata Internasional
Dosen: Moch Fahmi Firmasyah, S.Sy. M.A.
Penyusun: Moh. Raihan & Sufyan Tsaury
UIN Siber Syekh Nurjati Cirebon 2026
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Theme Colors: Deep Royal Navy & Imperial Gold
    BG_NAVY = RGBColor(4, 18, 40)       # #041228
    CARD_BG = RGBColor(10, 32, 66)      # #0A2042
    CARD_HOVER = RGBColor(14, 44, 90)   # #0E2C5A
    GOLD = RGBColor(229, 193, 88)       # #E5C158
    GOLD_LIGHT = RGBColor(254, 240, 138)# #FEF08A
    BLUE_LIGHT = RGBColor(96, 165, 250) # #60A5FA
    WHITE = RGBColor(255, 255, 255)
    LIGHT_GRAY = RGBColor(203, 213, 225)
    MUTED_GRAY = RGBColor(148, 163, 184)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_NAVY
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="HUKUM PERDATA INTERNASIONAL • MOH. RAIHAN & SUFYAN TSAURY", badge_num=""):
        # Header text
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10.5), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        
        # Category Tag
        p0 = tf.paragraphs[0]
        p0.text = category.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = GOLD
        p0.font.name = "Arial"

        # Title
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = WHITE
        p1.font.name = "Georgia"

        # Badge Number on top right
        if badge_num:
            num_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.35), Inches(1.0), Inches(0.9))
            ntf = num_box.text_frame
            np = ntf.paragraphs[0]
            np.text = badge_num
            np.font.size = Pt(36)
            np.font.bold = True
            np.font.color.rgb = RGBColor(60, 85, 125)
            np.font.name = "Georgia"
            np.alignment = PP_ALIGN.RIGHT

    # ==========================================
    # SLIDE 1: COVER
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Frame Border
    border = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(11.733), Inches(6.5))
    border.fill.solid()
    border.fill.fore_color.rgb = CARD_BG
    border.line.color.rgb = GOLD
    border.line.width = Pt(1.5)

    # Logo
    logo_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "logo-uinssc.png")
    if os.path.exists(logo_path):
        s1.shapes.add_picture(logo_path, Inches(6.1), Inches(0.8), height=Inches(0.95))

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.9), Inches(10.9), Inches(2.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "TUGAS MATA KULIAH HUKUM PERDATA INTERNASIONAL • 2026"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GOLD
    p.alignment = PP_ALIGN.CENTER

    p = tf1.add_paragraph()
    p.text = "SEJARAH & PERKEMBANGAN\nHUKUM PERDATA INTERNASIONAL"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Georgia"
    p.alignment = PP_ALIGN.CENTER

    p = tf1.add_paragraph()
    p.text = "Evolusi Pemikiran dari Teori Statuta Italia & Prancis, Mazhab Belanda & Anglo-Amerika, hingga Transformasi Regulasi Mandiri di Indonesia"
    p.font.size = Pt(13)
    p.font.color.rgb = LIGHT_GRAY
    p.alignment = PP_ALIGN.CENTER

    # Meta Cards
    box_dos = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.8), Inches(4.8), Inches(1.5))
    box_dos.fill.solid()
    box_dos.fill.fore_color.rgb = BG_NAVY
    box_dos.line.color.rgb = GOLD
    box_dos.line.width = Pt(1)
    tf_dos = box_dos.text_frame
    tf_dos.word_wrap = True
    p = tf_dos.paragraphs[0]
    p.text = "DOSEN PENGAMPU:"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = GOLD
    p = tf_dos.add_paragraph()
    p.text = "Moch Fahmi Firmasyah, S.Sy. M.A."
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p = tf_dos.add_paragraph()
    p.text = "Dosen Pengampu Hukum Perdata Internasional"
    p.font.size = Pt(9)
    p.font.color.rgb = MUTED_GRAY

    box_tim = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.8), Inches(5.0), Inches(1.5))
    box_tim.fill.solid()
    box_tim.fill.fore_color.rgb = BG_NAVY
    box_tim.line.color.rgb = GOLD
    box_tim.line.width = Pt(1)
    tf_tim = box_tim.text_frame
    tf_tim.word_wrap = True
    p = tf_tim.paragraphs[0]
    p.text = "DISUSUN OLEH:"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = GOLD
    p = tf_tim.add_paragraph()
    p.text = "Moh. Raihan [2530311065]\nSufyan Tsaury [2530311086]"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p = tf_tim.add_paragraph()
    p.text = "Prodi Hukum Keluarga • Fakultas Syari'ah • UINSSC"
    p.font.size = Pt(9)
    p.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 2: LATAR BELAKANG
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Latar Belakang & Urgensi Globalisasi", "BAB I PENDAHULUAN • LATAR BELAKANG", "02")

    factors = [
        ("Perpindahan Penduduk", "Mobilitas manusia antarnegara semakin intensif untuk bekerja, studi, dan migrasi domisili permanen."),
        ("Perkawinan Campuran", "Menimbulkan persoalan syarat sah nikah, harta perkawinan, status anak, dan perceraian."),
        ("Bisnis Transnasional", "Kontrak dagang internasional, e-commerce lintas negara, dan pengadaan ekspor-impor."),
        ("Kepemilikan Aset LN", "Sengketa kepemilikan aset properti atau rekening bisnis yang berada di luar negeri."),
        ("Unsur Asing (Foreign Element)", "Menuntut kepastian sistem hukum mana yang berhak mengadili dan dipakai menyelesaikan perkara.")
    ]
    for i, (title, desc) in enumerate(factors):
        x = Inches(0.8 + (i * 2.4))
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.9), Inches(2.2), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = GOLD
        card.line.width = Pt(1)
        ctf = card.text_frame
        ctf.word_wrap = True
        
        p = ctf.paragraphs[0]
        p.text = f"FAKTOR 0{i+1}"
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = GOLD
        p.alignment = PP_ALIGN.CENTER
        
        p = ctf.add_paragraph()
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        
        p = ctf.add_paragraph()
        p.text = desc
        p.font.size = Pt(9.5)
        p.font.color.rgb = LIGHT_GRAY
        p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 3: PENGERTIAN SINGKAT HPI
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Pengertian Singkat HPI & Unsur Asing", "BAB II PEMBAHASAN • KONSEP DASAR", "03")

    left_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(6.8), Inches(5.0))
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = CARD_BG
    left_box.line.color.rgb = GOLD
    left_box.line.width = Pt(1)
    ltf = left_box.text_frame
    ltf.word_wrap = True

    p = ltf.paragraphs[0]
    p.text = '"HPI bukan kodifikasi perdata dunia, melainkan aturan penunjuk nasional (referral rules) layaknya rambu lalu lintas hukum untuk menemukan lex causae."'
    p.font.size = Pt(11)
    p.font.italic = True
    p.font.color.rgb = GOLD_LIGHT

    p = ltf.add_paragraph()
    p.text = "\n5 TITIK TAUT / FAKTOR PENGAIT UNSUR ASING:"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = GOLD

    points3 = [
        "1. Kewarganegaraan: WNI menikah dengan WNA Jepang (lex patriae).",
        "2. Domisili: Penggugat bertempat tinggal di Singapura menggugat di RI.",
        "3. Tempat Peristiwa: Perbuatan melawan hukum di LN (locus delicti).",
        "4. Badan Hukum: Kontrak antara PT Indonesia dengan korporasi AS.",
        "5. Letak Objek: Sengketa kepemilikan tanah di luar negeri (lex rei sitae)."
    ]
    for pt in points3:
        p = ltf.add_paragraph()
        p.text = pt
        p.font.size = Pt(9.5)
        p.font.color.rgb = LIGHT_GRAY

    right_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.9), Inches(1.8), Inches(4.6), Inches(5.0))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = CARD_BG
    right_box.line.color.rgb = BLUE_LIGHT
    right_box.line.width = Pt(1)
    rtf = right_box.text_frame
    rtf.word_wrap = True
    p = rtf.paragraphs[0]
    p.text = "3 CONTOH KASUS NYATA:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = BLUE_LIGHT

    cases = [
        ("Perkawinan Campuran di Australia:", "WNI menikah dengan WNA Inggris di Australia; hukum mana yang mengatur keabsahan dan pemisahan harta?"),
        ("Kontrak Mesin Jerman:", "Pabrik RI memesan mesin ke Jerman namun barang rusak di laut; hukum mana yang mengatur ganti rugi?"),
        ("Waris Lintas Batas:", "Pewaris wafat meninggalkan properti di berbagai negara; hukum asal pewaris atau letak tanah yang berlaku?")
    ]
    for c_tit, c_desc in cases:
        p = rtf.add_paragraph()
        p.text = f"\n• {c_tit}"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p = rtf.add_paragraph()
        p.text = c_desc
        p.font.size = Pt(9)
        p.font.color.rgb = LIGHT_GRAY

    # ==========================================
    # SLIDE 4: AWAL PERKEMBANGAN HPI
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Awal Perkembangan HPI: Evolusi Sebab-Akibat", "BAB II PEMBAHASAN • SEJARAH AWAL", "04")

    steps = [
        ("Perbedaan Hukum", "Komunitas kuno hidup terisolasi dengan adat hukum lokalnya sendiri tanpa benturan."),
        ("Peningkatan Mobilitas", "Jalur maritim dan perdagangan antarkota di Eropa terbuka mempertemukan warga."),
        ("Hubungan Unsur Asing", "Pedagang lintas batas melakukan transaksi bisnis, pernikahan, dan kontrak."),
        ("Benturan Antarsistem", "Hakim menghadapi dilema: tunduk hukum kota asal atau hukum tempat transaksi?"),
        ("Lahirnya Teori HPI", "Kebutuhan keadilan melahirkan doktrin Teori Statuta hingga mazhab HPI modern.")
    ]
    for i, (title, desc) in enumerate(steps):
        x = Inches(0.8 + (i * 2.4))
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.9), Inches(2.2), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = GOLD
        card.line.width = Pt(1)
        ctf = card.text_frame
        ctf.word_wrap = True
        
        p = ctf.paragraphs[0]
        p.text = f"TAHAP 0{i+1}"
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = GOLD
        p.alignment = PP_ALIGN.CENTER
        
        p = ctf.add_paragraph()
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        
        p = ctf.add_paragraph()
        p.text = desc
        p.font.size = Pt(9.5)
        p.font.color.rgb = LIGHT_GRAY
        p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 5: STATUTA ITALIA
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Trikotomi Statuta Bartolus de Saxoferrato (Abad 13-15)", "BAB II PEMBAHASAN • MASA STATUTA ITALIA", "05")

    statutas = [
        ("Statuta Personalia", "Ekstrateritorial", "Mengatur status personal, kapasitas hukum, kedewasaan, dan keluarga.", "Melekat pada pribadi ke mana pun bepergian melintasi batas kota.", "Lex Personalia / Lex Patriae"),
        ("Statuta Realia", "Teritorial Murni", "Mengatur status kebendaan (khususnya tanah dan properti tidak bergerak).", "Hukum yang berlaku mutlak adalah hukum di tempat benda terletak (lex rei sitae).", "Lex Rei Sitae / Lex Situs"),
        ("Statuta Mixta", "Teritorial Terbatas", "Mengatur bentuk formalitas perbuatan hukum, akta kontrak, dan PMH.", "Hukum mengacu mutlak pada tempat di mana perbuatan itu dilaksanakan.", "Locus Regit Actum / Lex Loci Actus")
    ]
    for i, (name, nature, scope, rule, latin) in enumerate(statutas):
        x = Inches(0.8 + (i * 4.0))
        box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.9), Inches(3.7), Inches(4.8))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = GOLD
        box.line.width = Pt(1)
        btf = box.text_frame
        btf.word_wrap = True

        p = btf.paragraphs[0]
        p.text = f"KATEGORI 0{i+1}"
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = GOLD

        p = btf.add_paragraph()
        p.text = name
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.name = "Georgia"

        p = btf.add_paragraph()
        p.text = f"Sifat: {nature}"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = BLUE_LIGHT

        p = btf.add_paragraph()
        p.text = f"\nFokus Pengaturan:\n{scope}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = LIGHT_GRAY

        p = btf.add_paragraph()
        p.text = f"\nPrinsip Keberlakuan:\n{rule}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = LIGHT_GRAY

        p = btf.add_paragraph()
        p.text = f"\nAsas: {latin}"
        p.font.size = Pt(9)
        p.font.italic = True
        p.font.color.rgb = GOLD_LIGHT

    # ==========================================
    # SLIDE 6: STATUTA PRANCIS
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Statuta Prancis (Abad 16): Otonomi vs Kedaulatan", "BAB II PEMBAHASAN • STATUTA PRANCIS", "06")

    c1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = GOLD
    c1.line.width = Pt(1.5)
    c1tf = c1.text_frame
    c1tf.word_wrap = True

    p = c1tf.paragraphs[0]
    p.text = "Charles Dumoulin (1500–1566)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Georgia"

    p = c1tf.add_paragraph()
    p.text = "PENDUKUNG OTONOMI PARA PIHAK (PARTY AUTONOMY)"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = GOLD

    p = c1tf.add_paragraph()
    p.text = "\n• Gagasan Pokok:\nMenolak pembagian kaku ala Bartolus yang membatasi perdagangan antarprovinsi Prancis. Dumoulin mencetuskan asas Kebebasan Memilih Hukum (Party Autonomy)."
    p.font.size = Pt(10)
    p.font.color.rgb = LIGHT_GRAY

    p = c1tf.add_paragraph()
    p.text = "\n• Penerapan:\nDalam suatu kontrak bisnis, para pihak secara bebas dan tersirat berhak menyepakati hukum mana yang mengatur perjanjian mereka."
    p.font.size = Pt(10)
    p.font.color.rgb = LIGHT_GRAY

    p = c1tf.add_paragraph()
    p.text = "\n★ Pengaruh Modern: Meletakkan fondasi hukum kontrak komersial internasional modern."
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = GOLD_LIGHT

    c2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.9), Inches(5.6), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BLUE_LIGHT
    c2.line.width = Pt(1.5)
    c2tf = c2.text_frame
    c2tf.word_wrap = True

    p = c2tf.paragraphs[0]
    p.text = "Bertrand d'Argentré (1519–1590)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Georgia"

    p = c2tf.add_paragraph()
    p.text = "PENJAGA OTONOMI TERITORIAL FEODAL"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = BLUE_LIGHT

    p = c2tf.add_paragraph()
    p.text = "\n• Gagasan Pokok:\nMenolak pandangan Dumoulin yang dinilai terlalu liberal. Sebagai bangsawan Brittany, ia membendung hegemoni Paris dengan menegaskan kedaulatan wilayah tanah."
    p.font.size = Pt(10)
    p.font.color.rgb = LIGHT_GRAY

    p = c2tf.add_paragraph()
    p.text = "\n• Penerapan:\nPada prinsipnya SEMUA STATUTA ADALAH REALIA (teritorial). Pengecualian hanya diberikan secara sangat sempit bagi status pribadi seseorang."
    p.font.size = Pt(10)
    p.font.color.rgb = LIGHT_GRAY

    p = c2tf.add_paragraph()
    p.text = "\n★ Pengaruh Modern: Menegaskan batas-batas kedaulatan mutlak wilayah negara atas properti."
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = GOLD_LIGHT

    # ==========================================
    # SLIDE 7: BELANDA & ANGLO-AMERIKA
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Aliran Belanda & Anglo-Amerika", "BAB II PEMBAHASAN • MAZHAB BELANDA & COMMON LAW", "07")

    c_bel = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.8))
    c_bel.fill.solid()
    c_bel.fill.fore_color.rgb = CARD_BG
    c_bel.line.color.rgb = GOLD
    c_bel.line.width = Pt(1)
    t_bel = c_bel.text_frame
    t_bel.word_wrap = True

    p = t_bel.paragraphs[0]
    p.text = "Aliran Belanda (Abad ke-17)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Georgia"

    p = t_bel.add_paragraph()
    p.text = "Tokoh: Ulrik Huber, Johannes Voet, Paulus Voet"
    p.font.size = Pt(9.5)
    p.font.color.rgb = GOLD

    p = t_bel.add_paragraph()
    p.text = "\n1. Prinsip Teritorialitas Mutlak:\nPasca Perjanjian Westphalia 1648, hukum negara hanya berlaku mutlak mengikat di dalam batas batas teritorial wilayahnya sendiri."
    p.font.size = Pt(10)
    p.font.color.rgb = LIGHT_GRAY

    p = t_bel.add_paragraph()
    p.text = "\n2. Doktrin Comitas Gentium (Comity):\nPengadilan sukarela mengakui hukum asing atau hak yang diperoleh di luar negeri (vested rights) bukan karena paksaan, melainkan atas dasar kepatutan, kesopanan, dan demi kepentingan bersama para pedagang."
    p.font.size = Pt(10)
    p.font.color.rgb = LIGHT_GRAY

    c_ang = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.9), Inches(5.6), Inches(4.8))
    c_ang.fill.solid()
    c_ang.fill.fore_color.rgb = CARD_BG
    c_ang.line.color.rgb = BLUE_LIGHT
    c_ang.line.width = Pt(1)
    t_ang = c_ang.text_frame
    t_ang.word_wrap = True

    p = t_ang.paragraphs[0]
    p.text = "Aliran Anglo-Amerika (Abad ke-19)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Georgia"

    p = t_ang.add_paragraph()
    p.text = "Tokoh: Joseph Story (1834), A.V. Dicey, Joseph Beale"
    p.font.size = Pt(9.5)
    p.font.color.rgb = BLUE_LIGHT

    p = t_ang.add_paragraph()
    p.text = "\n1. Pendekatan Praktis Peradilan (Case Law):\nBerkembang secara induktif melalui preseden putusan pengadilan dalam naungan sistem Common Law dan federalisme Amerika Serikat."
    p.font.size = Pt(10)
    p.font.color.rgb = LIGHT_GRAY

    p = t_ang.add_paragraph()
    p.text = "\n2. Pemisahan Dua Pilar Utama:\n• Choice of Jurisdiction: Menguji apakah forum pengadilan berwenang mengadili.\n• Choice of Law: Menentukan hukum materiil yang memiliki keterikatan terdekat (most significant relationship)."
    p.font.size = Pt(10)
    p.font.color.rgb = LIGHT_GRAY

    # ==========================================
    # SLIDE 8: TABEL PERBANDINGAN MAZHAB
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Matriks Komparasi 4 Mazhab Besar Sejarah HPI", "BAB II PEMBAHASAN • ANALISIS KOMPARATIF", "08")

    rows = [
        ("Mazhab / Masa", "Tokoh Kunci", "Pokok Pemikiran & Ciri Khas", "Kontribusi bagi HPI Modern"),
        ("Statuta Italia\n(Abad 13–15)", "Accursius, Bartolus de Saxoferrato, Baldus", "Membedakan statuta menjadi personalia, realia, dan mixta. Perlindungan hak warga dari pemaksaan hukum kota lain.", "Peletak batu pertama doktrin lex personalia dan lex rei sitae."),
        ("Statuta Prancis\n(Abad ke-16)", "Charles Dumoulin, Bertrand d’Argentré", "Perdebatan otonomi kehendak para pihak vs dominasi kedaulatan tanah feodal (statuta realia).", "Melahirkan asas Party Autonomy dalam kontrak internasional."),
        ("Aliran Belanda\n(Abad ke-17)", "Ulrik Huber, Johannes & Paulus Voet", "Kedaulatan wilayah mutlak; pengakuan hukum asing berdasarkan asas kesopanan antarbangsa (Comitas Gentium).", "Menjelaskan rasio logis penerapan sukarela hukum asing tanpa kurangi martabat negara."),
        ("Anglo-Amerika\n(Abad 19–Modern)", "Joseph Story, A.V. Dicey, Joseph Beale", "Pendekatan berbasis peradilan (case law). Pemisahan tegas antara kewenangan forum dan pilihan hukum.", "Menghasilkan metode penyelesaian sengketa komersial & arbitrase modern.")
    ]
    table_shape = s8.shapes.add_table(5, 4, Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(2.4)
    table.columns[2].width = Inches(4.0)
    table.columns[3].width = Inches(3.133)

    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            cell.fill.solid()
            if r_idx == 0:
                cell.fill.fore_color.rgb = CARD_HOVER
                p = cell.text_frame.paragraphs[0]
                p.font.bold = True
                p.font.size = Pt(10.5)
                p.font.color.rgb = GOLD
            else:
                cell.fill.fore_color.rgb = CARD_BG
                p = cell.text_frame.paragraphs[0]
                p.font.size = Pt(9.5)
                p.font.color.rgb = WHITE if c_idx == 0 else LIGHT_GRAY
                if c_idx == 3:
                    p.font.color.rgb = GOLD_LIGHT

    # ==========================================
    # SLIDE 9: PERKEMBANGAN HPI DI INDONESIA
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Perkembangan HPI di Indonesia: Menuju Kodifikasi", "BAB II PEMBAHASAN • KONTEKS INDONESIA", "09")

    indo_cards = [
        ("Warisan Kolonial AB 1847", "Akar Historis Intergentiel Recht (Pasal 163 & 131 IS):\n\n• Pasal 16 AB: Asas personalitas (lex patriae), wewenang hukum WNI tetap melekat di mana pun berada.\n\n• Pasal 17 AB: Asas letak benda (lex rei sitae) bagi benda tidak bergerak.\n\n• Pasal 18 AB: Formalitas tindakan hukum (locus regit actum) dinilai menurut tempat pembuatan."),
        ("Regulasi Sektoral Tersebar", "Belum ada UU Kodifikasi HPI Mandiri; tercecer dalam berbagai aturan:\n\n• UU Perkawinan (UU 1/1974 jo UU 16/2019): Pengaturan perkawinan campuran (Pasal 57–62).\n\n• UU Adminduk (UU 24/2013): Pencatatan sipil peristiwa WNI di LN.\n\n• UU Arbitrase (UU 30/1999): Eksekusi putusan arbitrase internasional (New York Convention 1958)."),
        ("Urgensi RUU HPI Nasional", "Inisiatif BPHN & Kemenkumham (2023) untuk era ekonomi digital:\n\n★ Memberi kepastian hukum bagi WNI dan entitas bisnis nasional saat bermitra dengan pihak asing.\n\n★ Mengatur kejelasan pilihan hukum (choice of law) & forum (choice of court).\n\n★ Pedoman uji ketertiban umum (public policy) tanpa mencederai kepercayaan investasi global.")
    ]
    for i, (tit, desc) in enumerate(indo_cards):
        x = Inches(0.8 + (i * 4.0))
        box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.9), Inches(3.7), Inches(4.8))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = GOLD if i == 2 else RGBColor(70, 95, 135)
        box.line.width = Pt(1.5 if i == 2 else 1)
        btf = box.text_frame
        btf.word_wrap = True

        p = btf.paragraphs[0]
        p.text = f"PILAR 0{i+1}"
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = GOLD

        p = btf.add_paragraph()
        p.text = tit
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.name = "Georgia"

        p = btf.add_paragraph()
        p.text = f"\n{desc}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = LIGHT_GRAY

    # ==========================================
    # SLIDE 10: KESIMPULAN
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Kesimpulan & Menatap Masa Depan HPI", "BAB III PENUTUP • KESIMPULAN", "10")

    # Statement Box
    st_box = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.8), Inches(10.9), Inches(1.3))
    st_box.fill.solid()
    st_box.fill.fore_color.rgb = CARD_BG
    st_box.line.color.rgb = GOLD
    st_box.line.width = Pt(1.5)
    stf = st_box.text_frame
    stf.word_wrap = True
    p = stf.paragraphs[0]
    p.text = '"HPI berevolusi dari kebutuhan praktis kota dagang Italia menjadi instrumen strategis perlindungan kedaulatan dan kepastian bisnis global."'
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.italic = True
    p.font.color.rgb = GOLD_LIGHT
    p.font.name = "Georgia"
    p.alignment = PP_ALIGN.CENTER

    b1 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(3.4), Inches(5.2), Inches(3.3))
    b1.fill.solid()
    b1.fill.fore_color.rgb = CARD_BG
    b1.line.color.rgb = GOLD
    b1.line.width = Pt(1)
    b1tf = b1.text_frame
    b1tf.word_wrap = True
    p = b1tf.paragraphs[0]
    p.text = "Evolusi Doktrin Universal"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Georgia"
    p = b1tf.add_paragraph()
    p.text = "\nDari Teori Statuta, Comitas Gentium, hingga pemisahan yurisdiksi modern, setiap mazhab telah meletakkan fondasi esensial: perlindungan status orang (lex personalia), kepastian letak tanah (lex rei sitae), dan kebebasan berkontrak (party autonomy)."
    p.font.size = Pt(10)
    p.font.color.rgb = LIGHT_GRAY

    b2 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(3.4), Inches(5.2), Inches(3.3))
    b2.fill.solid()
    b2.fill.fore_color.rgb = CARD_BG
    b2.line.color.rgb = BLUE_LIGHT
    b2.line.width = Pt(1)
    b2tf = b2.text_frame
    b2tf.word_wrap = True
    p = b2tf.paragraphs[0]
    p.text = "Kebutuhan Strategis Indonesia"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Georgia"
    p = b2tf.add_paragraph()
    p.text = "\nMeninggalkan aturan peninggalan kolonial AB 1847 menuju pembentukan RUU HPI Nasional merupakan langkah mutlak guna mengawal kepentingan warga negara serta memperkokoh kepastian hukum transaksi ekonomi digital Indonesia di mata internasional."
    p.font.size = Pt(10)
    p.font.color.rgb = LIGHT_GRAY

    # Save output
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Sejarah_dan_Perkembangan_HPI.pptx")
    prs.save(output_path)
    print(f"PowerPoint berhasil dibuat: {output_path}")

if __name__ == "__main__":
    create_presentation()
